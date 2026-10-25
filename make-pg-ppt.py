# 토스페이먼츠 심사 제출용 결제경로 PPT.
# 그림은 운영 중인 dnalabs.kr 을 진짜 크롬 창으로 띄워 화면 전체를 찍은 것입니다 (shoot-pg.mjs).
#
# 토스 「홈페이지 결제경로 제작 가이드(정기결제용)」를 따릅니다.
#   · 반드시 PPT 형식 (카드사 심사에 맞게 고칠 수 있어야 해서) — PDF 는 만들지 않습니다
#   · 표지: 상호명 / 사업자번호 / URL / 테스트 계정
#   · 순서: ① 가맹점 정보 ② 하단 정보 ③ 환불 규정 ④ 로그인·회원가입 ⑤ 구매 과정 ⑥ 카드 입력창
import json, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

W, H = Inches(13.333), Inches(7.5)
남색 = RGBColor(0x0B, 0x1B, 0x33)
파랑 = RGBColor(0x2F, 0x6B, 0xFF)
흐림 = RGBColor(0x4E, 0x59, 0x68)
흰색 = RGBColor(0xFF, 0xFF, 0xFF)

기본 = os.environ['SHOTS']
장 = json.load(open(os.path.join(기본, 'pages.json'), encoding='utf-8'))

# 테스트 계정 — 있으면 표지에 적습니다. 비밀번호를 저장소에 남기지 않으려고
# 파일에 박지 않고 돌릴 때 환경변수로만 받습니다.
시험아이디 = os.environ.get('PG_TEST_ID', '').strip()
시험비번 = os.environ.get('PG_TEST_PW', '').strip()

p = Presentation()
p.slide_width, p.slide_height = W, H
빈장 = p.slide_layouts[6]

def 글(칸, 줄들, 크기, 색, 굵게=False, 정렬=PP_ALIGN.LEFT, 간격=None):
    tf = 칸.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(줄들, str): 줄들 = [줄들]
    for i, 내용 in enumerate(줄들):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = 정렬
        if 간격: para.space_after = Pt(간격)
        r = para.add_run(); r.text = 내용
        r.font.size = Pt(크기); r.font.color.rgb = 색; r.font.bold = 굵게
        r.font.name = 'Apple SD Gothic Neo'

# ── ① 표지 · 가맹점 정보 ──
s = p.slides.add_slide(빈장)
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = 남색
글(s.shapes.add_textbox(Inches(1.1), Inches(1.3), Inches(11), Inches(0.5)),
   '① 가맹점 정보', 16, RGBColor(0x9A, 0xC4, 0xFF), True)
글(s.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(11), Inches(1.0)),
   '결제 경로 안내 — Re:Service 정기결제', 36, 흰색, True)
줄 = [
    '(1) 상호명        디엔에이랩스 (DNA Labs)',
    '(2) 사업자번호    119-37-01707',
    '(3) URL             https://dnalabs.kr',
    '                       심사용 가입 주소  https://dnalabs.kr/signup?review=tosspay2026',
]
if 시험아이디 and 시험비번:
    줄 += [f'(4) Test ID        {시험아이디}', f'(5) Test PW       {시험비번}']
else:
    줄 += ['(4) 테스트 계정    별도 계정 없이 회원가입 화면에서 정기결제 카드 등록까지 진행됩니다']
글(s.shapes.add_textbox(Inches(1.1), Inches(3.2), Inches(11.5), Inches(3)),
   줄, 18, RGBColor(0xE1, 0xEA, 0xF6), 간격=10)
글(s.shapes.add_textbox(Inches(1.1), Inches(6.5), Inches(11), Inches(0.5)),
   '모든 화면은 운영 중인 사이트를 실제 크롬 창으로 열어 주소창과 PC 시각이 함께 보이도록 찍었습니다.',
   12, RGBColor(0x9F, 0xB3, 0xCC))

# 맨 아래 Dock(작업표시줄)은 잘라 냅니다 — 보기에 어수선하고 심사에 필요 없습니다.
# 맨 위 막대(시계)는 남깁니다. 토스 가이드가 PC 시각이 보이게 하라고 합니다.
# DOCK_PT = 화면 아래에서 Dock 이 차지하는 높이(포인트), SCREEN_H_PT = 화면 높이(포인트).
# 둘 다 build-pg-ppt.sh 가 macOS 에 물어 넘겨 줍니다. 없으면 자르지 않습니다.
덜어낼 = float(os.environ.get('DOCK_PT') or 0)
화면높이 = float(os.environ.get('SCREEN_H_PT') or 0)
def 아래자르기(길):
    if not (덜어낼 > 0 and 화면높이 > 0): return 길
    그림 = Image.open(길)
    배율 = 그림.height / 화면높이
    끝 = int(round((화면높이 - 덜어낼) * 배율))
    if not (0 < 끝 < 그림.height): return 길
    새길 = 길[:-4] + '-잘림.png'
    그림.crop((0, 0, 그림.width, 끝)).save(새길)
    return 새길

for 쪽 in 장:
    s = p.slides.add_slide(빈장)
    글(s.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(8.5), Inches(0.5)),
       쪽['제목'], 22, 남색, True)
    글(s.shapes.add_textbox(Inches(8.9), Inches(0.33), Inches(3.95), Inches(0.4)),
       쪽['주소'], 11, 파랑, 정렬=PP_ALIGN.RIGHT)
    글(s.shapes.add_textbox(Inches(0.5), Inches(0.78), Inches(12.33), Inches(0.5)),
       쪽['설명'], 12, 흐림)

    길 = 아래자르기(os.path.join(기본, 쪽['파일'] + '.png'))
    iw, ih = Image.open(길).size
    칸w, 칸h = Inches(12.33), Inches(5.95)
    배 = min(칸w / iw, 칸h / ih)
    w, h = Emu(int(iw * 배)), Emu(int(ih * 배))
    그림 = s.shapes.add_picture(길, Emu(int((W - w) / 2)), Inches(1.38), w, h)
    그림.line.color.rgb = RGBColor(0xD5, 0xDB, 0xE3); 그림.line.width = Pt(0.75)

out = os.environ['OUT']
p.save(out)
print('만들었습니다 →', out, os.path.getsize(out) // 1024, 'KB ·', len(장) + 1, '장')
